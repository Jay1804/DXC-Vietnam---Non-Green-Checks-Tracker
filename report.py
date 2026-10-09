"""
DXC Vietnam - Non Green Checks report (direct database version).

Flow: connect to MySQL -> pull non-green checks closed yesterday (or Fri-Sun when
run on a Monday) -> rename case_flex_fieldN columns using the client's field
definitions -> keep the final columns -> build an Excel workbook -> email it.
The email is always sent, even when there are no records.

Usage:
    python report.py                    # run for the dynamic date window and email it
    python report.py --dry-run          # build the workbook locally, do not email
    python report.py --today 2026-10-12 # pretend today is this date (testing)
"""

import argparse
import io
import os
import smtplib
import sys
from datetime import date, datetime, timedelta
from email.message import EmailMessage

import openpyxl
import pymysql
import pymysql.cursors
from dotenv import load_dotenv
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

load_dotenv()

CLIENT_NAME_LIKE = "%DXC Vietnam%"
REPORT_TITLE = "DXC Vietnam - Non Green Checks Report"

# Final report layout: (column in query result, header in report)
FINAL_COLUMNS = [
    ("candidate_name", "candidate_name"),
    ("company_name", "company_name"),
    ("case_ars_no", "case_ars_no"),
    ("process_name", "process_name"),
    ("location", "location"),
    ("case_received_date", "case_received_date"),
    ("Recruiter Name", "Recruiter Name"),
    ("Case initiator Email Id", "Case initiator Email Id"),
    ("check_name", "check_name"),
    ("go_ahead_date", "go_ahead_date"),
    ("check_closure_date", "check_closure_date"),
    ("check_severity", "check_severity"),
    ("closure_comments", "closure_comments"),
    ("latest_report_sent", "latest_report_sent"),
    ("candidate_email_id", "candidate_email_id"),
    ("check_created_on", "check_created_on"),
    ("institute_name_derived", "institute_name_derived"),
    ("check_unique_name", "check_unique_name"),
]

# --------------------------------------------------------------------------
# Query 1 - non-green checks (MySQL syntax; closure window is parameterised)
# --------------------------------------------------------------------------
FLEX_FIELD_COUNT = 30
_FLEX_SELECT = ",\n    ".join(f"ecff.case_flex_field{i}" for i in range(1, FLEX_FIELD_COUNT + 1))

QUERY_CHECKS = f"""
SELECT
    ecc.case_check_id,
    ecc.case_id,
    ecm.case_ars_no,
    emc.company_name,
    TRIM(CONCAT_WS(' ',
        NULLIF(ecc1.first_name, ''),
        NULLIF(ecc1.middle_name, ''),
        NULLIF(ecc1.last_name, ''))) AS candidate_name,
    ecp.process_name,
    emcl.office_name AS `location`,
    ecm.received_date AS case_received_date,
    ecm.created_date AS case_created_date,
    ecm.case_expected_closure_date AS case_due_date,
    ecm.last_insuff_date,
    ecm.last_insuff_fulfill_date,
    {_FLEX_SELECT},
    fn_case_status(ecm.case_status) AS case_status,
    ecc.check_name AS check_name,
    CASE WHEN ecc.family_id = 5 THEN emc1.company_name ELSE '' END AS company_name_derived,
    CASE WHEN ecc.family_id = 4 THEN emei.institute_name ELSE '' END AS institute_name_derived,
    ecc.check_created_on,
    ecc.go_ahead_date,
    ecc.check_expected_closure,
    ecc.insuff_date,
    ecc.insuff_remarks,
    ecc.insuff_fulfill_date,
    TRIM(CONCAT_WS(' ',
        NULLIF(eud.user_first_name, ''),
        NULLIF(eud.user_last_name, ''))) AS check_verifier,
    fn_check_status(ecc.check_status) AS check_status,
    ecc.check_severity,
    ecc.closure_comments,
    ec.check_name AS check_unique_name,
    ecc.check_closure_date,
    (
        SELECT ecr.report_sent_on
        FROM ec_case_reports ecr
        WHERE ecr.case_id = ecm.case_id
          AND ecr.report_status = 5
        ORDER BY ecr.case_report_id DESC
        LIMIT 1
    ) AS latest_report_sent,
    emei1.institute_name AS university_name,
    emei.allocation_field AS inst_allocation_field,
    emei1.allocation_field AS univ_allocation_field,
    (
        SELECT GROUP_CONCAT(eccd.stated_data ORDER BY eccd.stated_data SEPARATOR ', ')
        FROM ec_case_check_data eccd
        WHERE eccd.case_check_id = ecc.case_check_id
          AND eccd.field_id IN (1095, 1225, 1842, 2741, 2752)
    ) AS year_of_passing,
    ecc1.email_address AS candidate_email_id
FROM ec_case_master ecm
INNER JOIN ec_case_checks ecc ON ecm.case_id = ecc.case_id
LEFT JOIN ec_case_fields ecff ON ecm.case_id = ecff.case_id
LEFT JOIN ec_master_company emc ON emc.company_id = ecm.client_id
LEFT JOIN ec_client_process ecp ON ecm.process_id = ecp.process_id
LEFT JOIN ec_case_candidates ecc1 ON ecc1.candidate_id = ecm.candidate_id
LEFT JOIN ec_master_company_locations emcl ON ecm.client_office_id = emcl.office_id
LEFT JOIN ec_user_details eud ON ecc.check_verifier = eud.user_id
LEFT JOIN ec_case_check_verification_source eccvs ON ecc.case_check_id = eccvs.case_check_id
LEFT JOIN ec_master_company emc1 ON eccvs.org_id = emc1.company_id
LEFT JOIN ec_master_educational_institute emei ON eccvs.org_id = emei.institute_id
LEFT JOIN ec_master_educational_institute emei1 ON emei.university_id = emei1.institute_id
LEFT JOIN ec_checks ec ON ecc.check_id = ec.check_id
WHERE emc.company_name LIKE %(client)s
  AND ecc.check_closure_date >= %(start)s
  AND ecc.check_closure_date < %(end)s
  AND ecc.check_severity <> 'Green'
  AND ecc.check_severity <> 'Green_Approved'
  AND NULLIF(TRIM(ecc.check_severity), '') IS NOT NULL
"""

# --------------------------------------------------------------------------
# Query 2 - client case-field definitions (field_id N <-> case_flex_fieldN)
# --------------------------------------------------------------------------
QUERY_FIELDS = """
SELECT eccf.client_id,
       client_external_id,
       company_name AS 'Client Name',
       field_id,
       field_name,
       IF(`status`=0,'Passive','Active') AS 'Status',
       (SELECT GROUP_CONCAT(item_name) FROM ec_master_custom_data emcd
         WHERE item_status=1 AND eccf.FIELD_MASTER=emcd.MASTER_id) AS 'Drop_down_value',
       IF(IS_MANDATORY=0,'No','Yes') AS 'Mandatory Status'
FROM ec_client_case_fields eccf
LEFT JOIN ec_client ec ON eccf.client_id=ec.client_id
INNER JOIN ec_master_company emc ON eccf.client_id=emc.company_id
WHERE emc.company_name LIKE %(client)s
"""


def get_closure_window(today=None):
    """Return (start, end) dates for check_closure_date: start <= date < end.

    Normally yesterday. On a Monday: Friday, Saturday and Sunday.
    """
    today = today or date.today()
    days_back = 3 if today.weekday() == 0 else 1
    return today - timedelta(days=days_back), today


def connect():
    return pymysql.connect(
        host=os.environ["DB_HOST"],
        port=int(os.environ.get("DB_PORT", 3306)),
        user=os.environ["DB_USER"],
        password=os.environ["DB_PASSWORD"],
        database=os.environ["DB_NAME"],
        cursorclass=pymysql.cursors.DictCursor,
        charset="utf8mb4",
        connect_timeout=30,
        read_timeout=300,
    )


def fetch_data(start, end):
    """Run both queries; return (check rows, {flex column name: field_name})."""
    params = {"client": CLIENT_NAME_LIKE}
    with connect() as conn, conn.cursor() as cur:
        cur.execute(QUERY_CHECKS, {**params, "start": start, "end": end})
        rows = cur.fetchall()
        cur.execute(QUERY_FIELDS, params)
        fields = cur.fetchall()

    # Prefer Active definitions when a field_id is defined more than once.
    flex_map = {}
    for f in sorted(fields, key=lambda f: f["Status"] != "Active", reverse=True):
        flex_map[f"case_flex_field{f['field_id']}"] = f["field_name"]
    return rows, flex_map


def build_final_rows(rows, flex_map):
    """Rename flex columns to their field names and project to FINAL_COLUMNS."""
    out = []
    for r in rows:
        renamed = {flex_map.get(k, k): v for k, v in r.items()}
        out.append([_clean(renamed.get(src)) for src, _ in FINAL_COLUMNS])
    flex_names = {"Recruiter Name", "Case initiator Email Id"}
    missing = [src for src, _ in FINAL_COLUMNS
               if src in flex_names and src not in flex_map.values()]
    return [h for _, h in FINAL_COLUMNS], out, missing


def _clean(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    return v


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
BAND_FILL = PatternFill("solid", fgColor="F2F2F2")
SEVERITY_FILLS = {
    "Red": PatternFill("solid", fgColor="F8CBAD"),
    "Amber": PatternFill("solid", fgColor="FFE699"),
}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_workbook(headers, rows, start, end):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Final Report"
    last = get_column_letter(len(headers))

    last_day = end - timedelta(days=1)
    window = start.isoformat() if start == last_day else f"{start.isoformat()} to {last_day.isoformat()}"

    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = REPORT_TITLE
    ws["A1"].font = Font(bold=True, size=14, color="1F3864")
    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = (f"Check Closure Date: {window}    |    Records: {len(rows)}    |    "
                f"Generated: {datetime.now():%Y-%m-%d %H:%M}")
    ws["A2"].font = Font(italic=True, size=9, color="595959")

    hr = 4
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.fill, cell.border = HEADER_FILL, BORDER
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sev_idx = headers.index("check_severity")
    for i, row in enumerate(rows):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=hr + 1 + i, column=c, value=v)
            cell.font, cell.border = Font(size=10), BORDER
            cell.alignment = Alignment(vertical="center")
            if i % 2:
                cell.fill = BAND_FILL
        sev_fill = SEVERITY_FILLS.get(row[sev_idx])
        if sev_fill:
            ws.cell(row=hr + 1 + i, column=sev_idx + 1).fill = sev_fill

    for c, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[c - 1])) for r in rows])
        ws.column_dimensions[get_column_letter(c)].width = min(width + 3, 40)

    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    ws.auto_filter.ref = f"A{hr}:{last}{hr + max(len(rows), 1)}"
    if not rows:
        ws.merge_cells(f"A{hr + 1}:{last}{hr + 1}")
        note = ws.cell(row=hr + 1, column=1, value="No non-green checks found for the closure date(s) above.")
        note.font = Font(italic=True, color="808080", size=10)
        note.alignment = Alignment(horizontal="center")
    return wb, window


def workbook_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# Email
# --------------------------------------------------------------------------
def _recipients(var):
    return [a.strip() for a in os.environ.get(var, "").replace(";", ",").split(",") if a.strip()]


def send_email(window, row_count, filename, xlsx):
    to, cc = _recipients("REPORT_TO"), _recipients("REPORT_CC")
    if not to:
        raise RuntimeError("REPORT_TO is not set in .env (comma-separated recipient list).")
    sender = os.environ["EMAIL_SENDER"]

    msg = EmailMessage()
    msg["From"] = sender
    msg["To"] = ", ".join(to)
    if cc:
        msg["Cc"] = ", ".join(cc)
    msg["Subject"] = f"{REPORT_TITLE} - {window}"
    if row_count:
        msg.set_content(f"Hi,\n\nPlease find attached the {REPORT_TITLE} for check closure date "
                        f"{window}.\nTotal non-green checks: {row_count}.\n\nRegards,\nAuthbridge")
    else:
        msg.set_content(f"Hi,\n\nNo non-green checks were found for check closure date {window}.\n"
                        f"An empty report is attached for reference.\n\nRegards,\nAuthbridge")
    msg.add_attachment(xlsx, maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       filename=filename)

    with smtplib.SMTP(os.environ["SMTP_HOST"], int(os.environ.get("SMTP_PORT", 587)), timeout=60) as s:
        s.starttls()
        s.login(sender, os.environ["EMAIL_PASSWORD"])
        s.send_message(msg, to_addrs=to + cc)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dry-run", action="store_true", help="save the workbook locally, do not send email")
    ap.add_argument("--today", help="override today's date (YYYY-MM-DD) for testing")
    args = ap.parse_args()

    today = date.fromisoformat(args.today) if args.today else date.today()
    start, end = get_closure_window(today)
    print(f"Closure window: {start} <= check_closure_date < {end}")

    rows, flex_map = fetch_data(start, end)
    headers, final_rows, missing = build_final_rows(rows, flex_map)
    if missing:
        print(f"WARNING: columns not found in query/field mapping (left blank): {missing}")
    print(f"Rows: {len(final_rows)}")

    wb, window = build_workbook(headers, final_rows, start, end)
    xlsx = workbook_bytes(wb)
    filename = f"DXC_Vietnam_Non_Green_Checks_{start:%Y%m%d}_{(end - timedelta(days=1)):%Y%m%d}.xlsx"

    if args.dry_run:
        with open(filename, "wb") as f:
            f.write(xlsx)
        print(f"Dry run - saved {filename}, no email sent.")
        return
    send_email(window, len(final_rows), filename, xlsx)
    print("Email sent.")


if __name__ == "__main__":
    sys.exit(main())
