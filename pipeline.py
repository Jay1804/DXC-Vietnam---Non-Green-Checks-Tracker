"""
Shared pipeline logic for the DXC Vietnam Non-Green Tracker:
login -> pull MIS query -> filter severity -> select/rename columns ->
filter by closure date -> build a formatted Excel report.

Used by both download_dxc_report.py / build_final_report.py (CLI) and
app.py (Streamlit UI).
"""

import csv
import io
import zipfile
from datetime import datetime, timedelta

import openpyxl
import requests
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_URL = "https://mis.authbridge.com/export_query"
LOGIN_URL = f"{BASE_URL}/login.php"
PROCESS_URL = f"{BASE_URL}/process.php"
HEADERS = {"User-Agent": "Mozilla/5.0"}

HOSTNAME_ID = "4"              # Bridge Live
DATABASE = "checkpoint_live"   # Bridge Live database

QUERIES = {
    "daily": {
        "label": "Yesterday Done checks with other details as per client name - Bridge",
        "access_time": "Morning Data - Daily Dump",
        "csv_query": "909",
        "query_days_range": "0",
        "to_time": "23:59:59",
    },
    "tracker": {
        "label": "Client wise all cases with all checks with unique check name-Bridge On Received",
        "access_time": "Tracker - Query old",
        "csv_query": "1077",
        "query_days_range": "",
        "to_time": "23:59:59",
    },
}

# (source_column, output_label) - Company_name intentionally repeated, matching
# the layout requested for the Final Report.
SELECTED_COLUMNS = [
    ("Candidate_name", "Candidate_name"),
    ("Company_name", "Company_name"),
    ("case_ars_no", "case_ars_no"),
    ("Process_name", "Process_name"),
    ("location", "location"),
    ("case_received_date", "case_received_date"),
    ("case_flex_field1", "Candidate ID"),
    ("case_flex_field8", "Requisition ID"),
    ("Case_status", "Case_status"),
    ("check_created_on", "check_created_on"),
    ("Check_unique_name", "Check_unique_name"),
    ("check_status", "check_status"),
    ("check_severity", "check_severity"),
    ("CHECK_CLOSURE_DATE", "CHECK_CLOSURE_DATE"),
    ("candidate_email_id", "candidate_email_id"),
    ("institute_name", "institute_name"),
    ("Company_name", "Company_name"),
    ("closure_comments", "closure_comments"),
]

DATE_COLUMN_LABEL = "CHECK_CLOSURE_DATE"

HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11, name="Calibri")
TITLE_FONT = Font(color="1F3864", bold=True, size=16, name="Calibri")
SUBTITLE_FONT = Font(color="595959", italic=True, size=10, name="Calibri")
BODY_FONT = Font(size=10, name="Calibri")
BAND_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
SEVERITY_FILLS = {
    "Red": PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"),
    "Amber": PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid"),
}
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


class PipelineError(Exception):
    pass


def login(session, username, password):
    resp = session.post(
        LOGIN_URL,
        data={"username": username, "password": password, "login": "Login"},
        headers=HEADERS,
    )
    resp.raise_for_status()
    if "logout.php" not in resp.text and "searchForm" not in resp.text:
        raise PipelineError("Login failed - check credentials.")


def fetch_report_zip(session, query_key, from_dt, to_dt, client_name):
    query = QUERIES[query_key]
    date1 = from_dt.strftime("%Y-%m-%d 00:00:00")
    date2 = to_dt.strftime(f"%Y-%m-%d {query['to_time']}")

    resp = session.post(
        PROCESS_URL,
        data={
            "hostname": HOSTNAME_ID,
            "database": DATABASE,
            "access_time": query["access_time"],
            "csv_query": query["csv_query"],
            "query_days_range": query["query_days_range"],
            "date1": date1,
            "date2": date2,
            "client1": client_name,
        },
        headers=HEADERS,
        timeout=600,
    )
    resp.raise_for_status()

    content_type = resp.headers.get("Content-Type", "")
    if "zip" not in content_type and not resp.content.startswith(b"PK"):
        raise PipelineError(f"Unexpected response (not a zip file): {resp.text[:500]}")

    return resp.content


def extract_csv_rows(zip_bytes):
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        name = zf.namelist()[0]
        with zf.open(name) as f:
            text = io.TextIOWrapper(f, encoding="utf-8-sig", newline="")
            reader = csv.DictReader(text)
            fieldnames = reader.fieldnames
            rows = list(reader)
    return fieldnames, rows


def filter_severity(rows, severities):
    if not severities:
        return list(rows)
    return [r for r in rows if r.get("check_severity") in severities]


def select_and_rename(rows, columns=SELECTED_COLUMNS):
    out_headers = [label for _, label in columns]
    out_rows = [[r.get(src, "") for src, _ in columns] for r in rows]
    return out_headers, out_rows


def get_target_dates(today=None):
    today = today or datetime.now()
    if today.weekday() == 0:  # Monday -> Fri, Sat, Sun
        return {(today - timedelta(days=d)).strftime("%Y-%m-%d") for d in (1, 2, 3)}
    return {(today - timedelta(days=1)).strftime("%Y-%m-%d")}


def filter_by_closure_date(headers, rows, target_dates, date_label=DATE_COLUMN_LABEL):
    idx = headers.index(date_label)
    kept = []
    for row in rows:
        val = row[idx]
        date_part = str(val).split(" ")[0] if val else ""
        if date_part in target_dates:
            kept.append(row)
    return kept


def autosize_columns(ws, headers, start_row):
    for col_idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for row in ws.iter_rows(min_row=start_row, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 40)


def build_report_workbook(headers, rows, target_dates, title="DXC Vietnam - Non Green Checks Report",
                           severity_label="Amber, Red"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Final Report"

    n_cols = len(headers)
    last_col_letter = get_column_letter(n_cols)

    ws.merge_cells(f"A1:{last_col_letter}1")
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26

    date_label = " / ".join(sorted(target_dates)) if target_dates else "N/A"
    ws.merge_cells(f"A2:{last_col_letter}2")
    ws["A2"] = (f"Check Closure Date: {date_label}    |    Severity: {severity_label}    |    "
                f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    ws["A2"].font = SUBTITLE_FONT
    ws.row_dimensions[2].height = 16

    ws.row_dimensions[3].height = 6

    header_row = 4
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[header_row].height = 22

    severity_idx = headers.index("check_severity") if "check_severity" in headers else None

    for r_offset, row in enumerate(rows):
        r = header_row + 1 + r_offset
        band = BAND_FILL if r_offset % 2 == 1 else None
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col_idx, value=value)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="left", vertical="center")
            if band:
                cell.fill = band
        if severity_idx is not None:
            sev_val = row[severity_idx]
            sev_fill = SEVERITY_FILLS.get(sev_val)
            if sev_fill:
                ws.cell(row=r, column=severity_idx + 1).fill = sev_fill

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = (f"A{header_row}:{last_col_letter}{header_row + len(rows)}"
                           if rows else f"A{header_row}:{last_col_letter}{header_row}")

    autosize_columns(ws, headers, start_row=header_row)

    if not rows:
        ws.merge_cells(f"A{header_row + 1}:{last_col_letter}{header_row + 1}")
        note = ws.cell(row=header_row + 1, column=1, value="No records found for the selected date(s).")
        note.font = Font(italic=True, color="808080", size=10)
        note.alignment = Alignment(horizontal="center")

    return wb


def workbook_to_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
